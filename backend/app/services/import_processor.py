"""Handles CSV/Excel uploads with deduplication and conflict detection."""
import csv
import io
import uuid
from datetime import datetime
from typing import Any
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy import select
from app.models.physician import Physician
from app.models.import_batch import ImportBatch
from app.models.import_conflict import ImportConflict
from app.services.match_engine import MatchEngine
from app.services.completeness_scorer import CompletenessScorer

CONFLICT_FIELDS = ['credentials', 'specialty', 'institution_name', 'practice_type', 'city', 'state']

FIELD_MAPPING = {
    'first_name': 'first_name',
    'last_name': 'last_name', 
    'npi': 'npi',
    'credentials': 'credentials',
    'specialty': 'specialty',
    'subspecialty': 'subspecialty',
    'practice_type': 'practice_type',
    'institution_name': 'institution_name',
    'institution_type': 'institution_type',
    'city': 'city',
    'state': 'state',
    'region': 'region',
    'country': 'country',
    'years_in_practice': 'years_in_practice',
    'institutional_role': 'institutional_role',
}

class ImportProcessor:
    def __init__(self, db: AsyncSession):
        self.db = db
        self.match_engine = MatchEngine(db)
        self.completeness_scorer = CompletenessScorer(db)
    
    async def process_upload(self, file_content: bytes, filename: str, uploaded_by: str, team: str, description: str = "") -> ImportBatch:
        batch = ImportBatch(
            filename=filename,
            uploaded_by=uploaded_by,
            team=team,
            description=description,
            total_rows=0,
            new_records=0,
            updated_records=0,
            duplicate_records=0,
            error_records=0,
            status="processing"
        )
        self.db.add(batch)
        await self.db.flush()
        
        rows = self._parse_file(file_content, filename)
        batch.total_rows = len(rows)
        
        for row in rows:
            try:
                first_name = row.get('first_name', '').strip()
                last_name = row.get('last_name', '').strip()
                npi = row.get('npi', '').strip() or None
                state = row.get('state', '').strip() or None
                
                if not first_name or not last_name:
                    batch.error_records += 1
                    continue
                
                existing = await self.match_engine.find_match(npi, first_name, last_name, state)
                
                if not existing:
                    # New physician
                    physician = Physician(
                        first_name=first_name,
                        last_name=last_name,
                        record_status='imported',
                        source_channel='import',
                        source_detail=f"Import: {filename} by {uploaded_by} ({team})",
                        original_import_id=batch.id,
                    )
                    for csv_field, model_field in FIELD_MAPPING.items():
                        if csv_field in ('first_name', 'last_name'):
                            continue
                        val = row.get(csv_field, '').strip()
                        if val:
                            if model_field == 'years_in_practice':
                                try:
                                    setattr(physician, model_field, int(val))
                                except ValueError:
                                    pass
                            else:
                                setattr(physician, model_field, val)
                    self.db.add(physician)
                    await self.db.flush()
                    await self.completeness_scorer.compute(physician)
                    batch.new_records += 1
                    
                else:
                    # Check for conflicts
                    conflicts_found = False
                    for field in CONFLICT_FIELDS:
                        incoming_val = row.get(field, '').strip()
                        existing_val = getattr(existing, field, '') or ''
                        if incoming_val and existing_val and incoming_val.lower() != existing_val.lower():
                            conflict = ImportConflict(
                                import_batch_id=batch.id,
                                physician_id=existing.id,
                                conflict_type='attribute_mismatch',
                                field_name=field,
                                existing_value=existing_val,
                                incoming_value=incoming_val,
                                resolution='pending',
                            )
                            self.db.add(conflict)
                            conflicts_found = True
                    
                    if conflicts_found:
                        batch.duplicate_records += 1
                    else:
                        existing.updated_at = datetime.utcnow()
                        batch.updated_records += 1
                    
            except Exception:
                batch.error_records += 1
        
        batch.status = "completed"
        batch.completed_at = datetime.utcnow()
        await self.db.commit()
        return batch
    
    def _parse_file(self, content: bytes, filename: str) -> list[dict[str, str]]:
        if filename.endswith('.xlsx') or filename.endswith('.xls'):
            return self._parse_excel(content)
        return self._parse_csv(content)
    
    def _parse_csv(self, content: bytes) -> list[dict[str, str]]:
        text = content.decode('utf-8-sig')
        reader = csv.DictReader(io.StringIO(text))
        return [dict(row) for row in reader]
    
    def _parse_excel(self, content: bytes) -> list[dict[str, str]]:
        from openpyxl import load_workbook
        wb = load_workbook(filename=io.BytesIO(content), read_only=True)
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            return []
        headers = [str(h).strip().lower().replace(' ', '_') if h else '' for h in rows[0]]
        result = []
        for row in rows[1:]:
            record = {}
            for i, val in enumerate(row):
                if i < len(headers) and headers[i]:
                    record[headers[i]] = str(val).strip() if val is not None else ''
            result.append(record)
        return result
